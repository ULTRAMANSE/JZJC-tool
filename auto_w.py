# coding=utf-8
from docx import Document
from pypinyin import Style, lazy_pinyin
from PyQt5.Qt import QThread, QMutex
from PyQt5.QtCore import pyqtSignal
from read_word import read_head

lock = QMutex()


class Read(QThread):
	str_out = pyqtSignal(str)

	def __init__(self, docx_in=None, parent=None):
		super().__init__(parent)
		self.docx_in = docx_in
		self.doc = Document(self.docx_in)
		self.t_number = 5  # 计算
		self.temp_out = []  # 输出列表
		self.temp_num = 0
		self.test_style = {"功能性": "F", "性能效率": "P", "信息安全性": "IS", "兼容性": "Sc", "易用性": "Su", "可靠性": "Sr", "可维护": "Sm",
						   "可移植": "Sp"}

	def run(self):
		lock.lock()
		tb = self.doc.tables
		if "大纲" in self.docx_in:
			for index, tb_row in enumerate(tb[3].column_cells(3)):
				if index in (0, 1):
					continue
				tb[3].column_cells(2)[index].text = "".join(lazy_pinyin(tb_row.text, style=Style.FIRST_LETTER)).upper()
		else:
			self.re_style = read_head(self.docx_in)
			while self.t_number < 5 + len(self.re_style):
				temp_a = tb[self.t_number].cell(0, 2).text
				temp_b = tb[self.t_number].cell(0, 4).text
				if "用例标识" in temp_a and "用例名称" in temp_b:
					for index, tb_row in enumerate(tb[self.t_number].column_cells(4)):
						if index is 0:
							continue
						self.temp_out.append(
							"".join(lazy_pinyin(tb_row.text, style=Style.FIRST_LETTER)).upper() + "-" + self.test_style[
								self.re_style[self.t_number - 5]] + "-" + "".join(
								lazy_pinyin(tb[self.t_number].column_cells(0)[index].text,
											style=Style.FIRST_LETTER)).upper() + "-" + str(index).zfill(3))
				self.t_number += 1

			# 填写编号
			self.t_number = 5
			while self.t_number < 5 + len(self.re_style):
				temp_a = tb[self.t_number].cell(0, 2).text
				temp_b = tb[self.t_number].cell(0, 4).text
				if "用例标识" in temp_a and "用例名称" in temp_b:
					for index, tb_row in enumerate(tb[self.t_number].column_cells(4)):
						if index is 0:
							continue
						tb[self.t_number].column_cells(2)[index].text = self.temp_out[self.temp_num]
						self.temp_num += 1
				self.t_number += 1

		self.doc.save(self.docx_in)
		self.str_out.emit("执行完成")
		lock.unlock()
