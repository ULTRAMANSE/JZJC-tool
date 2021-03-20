# coding = utf-8
from docx import Document
from pypinyin import Style, lazy_pinyin


def read(docx_in):
	# print(docx_in)
	doc = Document(docx_in)
	tb = doc.tables
	example_name = []
	test_item = []
	t_number = 5
	c = None
	while t_number < 6:
		temp_a = tb[t_number].cell(0, 2).text
		temp_b = tb[t_number].cell(0, 4).text
		if "用例标识" in temp_a and "用例名称" in temp_b:
			for index, tb_row in enumerate(tb[t_number].column_cells(4)):
				if index is 0:
					continue
				example_name.append("".join(lazy_pinyin(tb_row.text, style=Style.FIRST_LETTER)).upper())
				test_item.append(
					"".join(lazy_pinyin(tb[t_number].column_cells(0)[index].text, style=Style.FIRST_LETTER)).upper())
				c = list(map(list, zip(test_item, example_name)))
		t_number += 1

	temp_out = []
	for i in range(len(c)):
		if i < 9:
			tstr = "00" + str(i + 1)
		elif i < 99:
			tstr = "0" + str(i + 1)
		else:
			tstr = str(i + 1)
		temp_out.append(c[i][0] + "-" + "F" + "-" + c[i][1] + "-" + tstr)

	t_number = 5
	while t_number < 6:
		temp_a = tb[t_number].cell(0, 2).text
		temp_b = tb[t_number].cell(0, 4).text
		if "用例标识" in temp_a and "用例名称" in temp_b:
			for index, tb_row in enumerate(tb[t_number].column_cells(4)):
				if index is 0:
					continue
				tb[t_number].column_cells(2)[index].text = temp_out[index - 1]
		t_number += 1
	doc.save(docx_in)
	return "执行完成"


if __name__ == '__main__':
	# read("./test1.docx")
	read(input("请复制文件完整路径："))
